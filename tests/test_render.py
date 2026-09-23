"""Report.render の E2E テスト(第7段)。

モデル → solver → PlacementPlan → openpyxl backend のパイプライン全体を
render 経由で通し、xlsx を読み戻して確認する。レイアウト計算の網羅は
test_solver.py が担うので、ここでは入力契約(dict / Sequence / 検証)と
配線の正しさだけを見る。
"""

import warnings
from typing import Annotated

import pytest
from openpyxl import load_workbook
from pydantic import BaseModel, Field, ValidationError

from schemaxl import PlacementPlan
from schemaxl.core.errors import LayoutError, SchemaxlWarning
from schemaxl.core.model import A4, Auto, Layout, Report, Table
from schemaxl.core.overflow import Shrink
from schemaxl.core.units import mm


class AllergyRow(BaseModel):
    child_name: Annotated[str, Field(max_length=5), Layout(header="児童名")]
    allergens: Annotated[list[str], Layout(header="アレルゲン")]


class AllergyReport(Report[AllergyRow]):
    page = A4(orientation="portrait", margin=mm(15))
    body = Table()  # bind はジェネリクスから推論される


def test_render_writes_dict_rows_to_xlsx(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    AllergyReport.render(
        [
            {"child_name": "山田", "allergens": ["卵", "乳"]},
            {"child_name": "佐藤", "allergens": ["小麦"]},
        ],
        str(path),
    )

    ws = load_workbook(str(path)).active
    assert ws.cell(row=1, column=1).value == "児童名"
    assert ws.cell(row=1, column=2).value == "アレルゲン"
    assert ws.cell(row=2, column=1).value == "山田"
    assert ws.cell(row=2, column=2).value == "卵、乳"  # list は join で 1 セルへ
    assert ws.cell(row=3, column=1).value == "佐藤"


def test_render_accepts_sequence_rows_in_field_declaration_order(tmp_path) -> None:
    # duckdb.fetchall() 由来のタプル列を想定(フィールド宣言順で位置マップ)。
    path = tmp_path / "report.xlsx"
    AllergyReport.render([("山田", ["卵", "乳"])], str(path))

    ws = load_workbook(str(path)).active
    assert ws.cell(row=2, column=1).value == "山田"
    assert ws.cell(row=2, column=2).value == "卵、乳"


def test_render_validates_rows_with_pydantic(tmp_path) -> None:
    # max_length=5 のデータ制約違反は書き出し前に ValidationError で止まる。
    with pytest.raises(ValidationError):
        AllergyReport.render(
            [{"child_name": "あいうえおかき", "allergens": []}],
            str(tmp_path / "report.xlsx"),
        )


def test_sequence_row_with_wrong_arity_is_rejected(tmp_path) -> None:
    with pytest.raises(ValueError, match="フィールド数"):
        AllergyReport.render([("山田",)], str(tmp_path / "report.xlsx"))


def test_explicit_bind_is_respected(tmp_path) -> None:
    class BoundReport(Report):  # ジェネリクスを書かず bind で明示するスタイル
        page = A4(margin=mm(15))
        body = Table(bind=AllergyRow)

    path = tmp_path / "report.xlsx"
    BoundReport.render([{"child_name": "山田", "allergens": []}], str(path))

    ws = load_workbook(str(path)).active
    assert ws.cell(row=2, column=1).value == "山田"


def test_render_without_bind_or_generic_raises_layout_error(tmp_path) -> None:
    class UnboundReport(Report):
        page = A4(margin=mm(15))
        body = Table()

    with pytest.raises(LayoutError):
        UnboundReport.render([], str(tmp_path / "report.xlsx"))


# --- 見切れ警告の扱い(strict) --------------------------------------------


class TightRow(BaseModel):
    label: Annotated[
        str,
        Layout(header="ラベル", width=Auto(min=mm(20)), overflow=Shrink(min_pt=8)),
    ]


class TightReport(Report[TightRow]):
    page = A4(orientation="portrait", margin=mm(15))
    body = Table()


CLIPPED = [{"label": "あ" * 30}]


def test_strict_render_refuses_to_write_a_clipped_report(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    with pytest.raises(LayoutError, match="strict=False"):
        TightReport.render(CLIPPED, str(path))
    # 見切れた帳票は書き出さない。
    assert not path.exists()


def test_non_strict_render_warns_and_still_writes(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    with pytest.warns(SchemaxlWarning, match="収まらない"):
        TightReport.render(CLIPPED, str(path), strict=False)
    assert path.exists()


def test_render_does_not_warn_when_everything_fits(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    with warnings.catch_warnings():
        warnings.simplefilter("error", SchemaxlWarning)
        AllergyReport.render([{"child_name": "山田", "allergens": ["卵"]}], str(path))
    assert path.exists()


# --- 拡張点(plan / measure / backend / sheet_name) -------------------------


class RecordingBackend:
    """受け取った plan と path を記録するだけの Backend。ファイルは書かない。"""

    def __init__(self) -> None:
        self.calls: list[tuple[PlacementPlan, object]] = []

    def write(self, plan: PlacementPlan, path: object) -> None:
        self.calls.append((plan, path))


def test_plan_returns_the_placement_plan_without_writing() -> None:
    plan = AllergyReport.plan([{"child_name": "山田", "allergens": ["卵", "乳"]}])
    assert isinstance(plan, PlacementPlan)
    values = {(c.row, c.col): c.value for c in plan.cells}
    assert values[(2, 1)] == "山田"
    assert values[(2, 2)] == "卵、乳"


def test_plan_keeps_warnings_instead_of_raising() -> None:
    # strict の判定は render の責務。plan は警告を載せて返すだけ。
    plan = TightReport.plan(CLIPPED)
    assert [w.kind for w in plan.warnings] == ["overflow"]


def test_render_hands_the_plan_to_the_given_backend(tmp_path) -> None:
    backend = RecordingBackend()
    path = tmp_path / "report.xlsx"
    AllergyReport.render([{"child_name": "山田", "allergens": []}], path, backend=backend)

    assert len(backend.calls) == 1
    plan, received_path = backend.calls[0]
    assert received_path == path
    assert plan == AllergyReport.plan([{"child_name": "山田", "allergens": []}])
    # 既定の openpyxl バックエンドは使われていない。
    assert not path.exists()


def test_strict_render_does_not_call_the_backend_on_warnings(tmp_path) -> None:
    backend = RecordingBackend()
    with pytest.raises(LayoutError):
        TightReport.render(CLIPPED, tmp_path / "report.xlsx", backend=backend)
    assert backend.calls == []


def test_injected_measure_drives_column_widths() -> None:
    rows = [{"child_name": "山田", "allergens": ["卵"]}]

    def wide(text: str, font_pt: float) -> float:
        return len(text) * font_pt * 3  # 既定の計測より 3 倍以上広く見積もる

    default_plan = AllergyReport.plan(rows)
    wide_plan = AllergyReport.plan(rows, measure=wide)
    assert wide_plan.column_widths[1] > default_plan.column_widths[1]


def test_sheet_name_is_carried_by_the_plan_and_written(tmp_path) -> None:
    rows = [{"child_name": "山田", "allergens": []}]
    assert AllergyReport.plan(rows, sheet_name="3年生").sheet_name == "3年生"

    path = tmp_path / "report.xlsx"
    AllergyReport.render(rows, path, sheet_name="3年生")
    assert load_workbook(str(path)).active.title == "3年生"


def test_render_accepts_pathlike(tmp_path) -> None:
    path = tmp_path / "report.xlsx"  # pathlib.Path のまま渡す
    AllergyReport.render([{"child_name": "山田", "allergens": []}], path)
    assert path.exists()
