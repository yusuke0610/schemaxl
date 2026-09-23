"""openpyxl バックエンドの薄いテスト(第6段)。

固定の PlacementPlan を書き出し、生成された xlsx を openpyxl で読み戻して
「solver の決定がファイルへ反映されているか」だけを確認する。網羅は不要。

このテストは backend 層のものなので openpyxl を import してよい
(core 層のテストは純 Python に保つ、という制約の対象外)。
"""

from dataclasses import replace

import pytest
from openpyxl import load_workbook

from schemaxl.backends.openpyxl_backend import write_xlsx
from schemaxl.core.plan import CellPlacement, CellRange, PageBreak, PageSetup, PlacementPlan
from schemaxl.core.units import mm


def _fixed_plan() -> PlacementPlan:
    return PlacementPlan(
        cells=[
            CellPlacement(row=1, col=1, value="児童名", font_pt=11),
            CellPlacement(row=1, col=2, value="アレルゲン", font_pt=11),
            CellPlacement(row=2, col=1, value="山田", font_pt=11, wrap=True),
            CellPlacement(row=2, col=2, value="卵、乳", font_pt=8),
            CellPlacement(row=3, col=1, value="佐藤", font_pt=11, wrap=True),
            CellPlacement(row=3, col=2, value="小麦", font_pt=8),
        ],
        page_breaks=[PageBreak(before_row=3)],
        column_widths={1: 12.0, 2: 20.0},
        row_heights={1: 14.0, 2: 13.2, 3: 13.2},
        header_rows=1,
        page=PageSetup(
            paper="A4",
            orientation="portrait",
            width_pt=mm(210).to_pt(),
            height_pt=mm(297).to_pt(),
            margin_top_pt=mm(15).to_pt(),
            margin_right_pt=mm(15).to_pt(),
            margin_bottom_pt=mm(15).to_pt(),
            margin_left_pt=mm(15).to_pt(),
        ),
        print_area=CellRange(first_row=1, first_col=1, last_row=3, last_col=2),
    )


def test_cell_values_are_written_and_read_back(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    write_xlsx(_fixed_plan(), str(path))

    ws = load_workbook(str(path)).active
    assert ws.cell(row=1, column=1).value == "児童名"
    assert ws.cell(row=1, column=2).value == "アレルゲン"
    assert ws.cell(row=2, column=2).value == "卵、乳"
    assert ws.cell(row=3, column=1).value == "佐藤"


def test_font_size_and_wrap_are_applied(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    write_xlsx(_fixed_plan(), str(path))

    ws = load_workbook(str(path)).active
    assert ws.cell(row=2, column=1).alignment.wrap_text is True  # Wrap セル
    assert ws.cell(row=2, column=2).font.size == 8  # Shrink 後のフォント


def test_repeat_header_becomes_print_title_rows(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    write_xlsx(_fixed_plan(), str(path))

    ws = load_workbook(str(path)).active
    # openpyxl は絶対参照へ正規化して読み戻す($1:$1)。
    assert ws.print_title_rows == "$1:$1"


def test_page_break_before_row_is_written_after_the_previous_row(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    write_xlsx(_fixed_plan(), str(path))

    ws = load_workbook(str(path)).active
    # before_row=3 → 行 2 の直後で改ページ(openpyxl の Break id=2)。
    break_ids = [brk.id for brk in ws.row_breaks.brk]
    assert 2 in break_ids


def test_column_widths_and_row_heights_are_written(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    write_xlsx(_fixed_plan(), str(path))

    ws = load_workbook(str(path)).active
    assert ws.column_dimensions["A"].width == 12.0
    assert ws.row_dimensions[2].height == 13.2


def test_paper_size_and_orientation_are_written(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    write_xlsx(_fixed_plan(), str(path))

    ws = load_workbook(str(path)).active
    assert ws.page_setup.paperSize == 9  # OOXML の A4
    assert ws.page_setup.orientation == "portrait"


def test_landscape_orientation_is_written(tmp_path) -> None:
    plan = _fixed_plan()
    assert plan.page is not None
    plan.page = replace(plan.page, orientation="landscape")
    path = tmp_path / "report.xlsx"
    write_xlsx(plan, str(path))

    ws = load_workbook(str(path)).active
    assert ws.page_setup.orientation == "landscape"


def test_margins_are_written_in_inches(tmp_path) -> None:
    """solver が pt で決めた余白が、OOXML の単位(インチ)へ変換されて載ること。"""
    path = tmp_path / "report.xlsx"
    write_xlsx(_fixed_plan(), str(path))

    ws = load_workbook(str(path)).active
    expected = 15 / 25.4  # 15mm をインチへ
    assert ws.page_margins.top == pytest.approx(expected)
    assert ws.page_margins.right == pytest.approx(expected)
    assert ws.page_margins.bottom == pytest.approx(expected)
    assert ws.page_margins.left == pytest.approx(expected)
    # ヘッダ / フッタ領域は明示的に 0(openpyxl 既定の 0.3 inch を残さない)。
    assert ws.page_margins.header == pytest.approx(0.0)
    assert ws.page_margins.footer == pytest.approx(0.0)


def test_print_area_is_written(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    write_xlsx(_fixed_plan(), str(path))

    ws = load_workbook(str(path)).active
    # ヘッダ 1 行 + データ 2 行 × 2 列。
    assert "$A$1:$B$3" in str(ws.print_area)


def test_plan_without_page_setup_leaves_defaults(tmp_path) -> None:
    """page が None の plan は書き出せる(ページ設定に触れない)。"""
    path = tmp_path / "report.xlsx"
    write_xlsx(PlacementPlan(cells=[CellPlacement(row=1, col=1, value="x")]), str(path))

    ws = load_workbook(str(path)).active
    assert ws.cell(row=1, column=1).value == "x"
    assert ws.print_area in ([], None, "")


def test_unknown_paper_is_rejected(tmp_path) -> None:
    plan = _fixed_plan()
    assert plan.page is not None
    plan.page = replace(plan.page, paper="B5")
    with pytest.raises(ValueError, match="B5"):
        write_xlsx(plan, str(tmp_path / "report.xlsx"))
